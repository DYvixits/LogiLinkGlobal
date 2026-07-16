from fastapi import FastAPI, APIRouter, HTTPException, Depends, Header
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field
from typing import List, Optional, Literal, Dict
from datetime import datetime, timedelta, timezone
from collections import defaultdict
import os
import random
import string
import io
import asyncio
import secrets
import base64
import pyotp
import qrcode
from openpyxl import Workbook
from passlib.context import CryptContext
from jose import JWTError, jwt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.graphics.barcode import code128
from reportlab.graphics.shapes import Drawing
from dotenv import load_dotenv

load_dotenv()

# --- Config & Auth Setup ---
SECRET_KEY = os.environ["JWT_SECRET"]
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = int(os.environ.get("ACCESS_TOKEN_EXPIRE_MINUTES", "720"))

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="/api/auth/login")

# MongoDB setup
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]
parcels_collection = db.parcels
users_collection = db.users
audit_collection = db.audit_logs
agencies_collection = db.agencies
settings_collection = db.settings
invoices_collection = db.invoices
apikeys_collection = db.api_keys

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

api_router = APIRouter(prefix="/api")

# --- Workflow statuses ---
# Linear lifecycle (order matters for the timeline / progress)
STATUS_FLOW = [
    "CREATED", "REGISTERED", "RECEIVED_AT_DEPOT", "CONTROLLED", "WEIGHED",
    "PACKED", "INVOICED", "PAID", "LOADED", "IN_TRANSIT", "IN_CUSTOMS",
    "ARRIVED", "AVAILABLE", "DELIVERED",
]
EXCEPTION_STATUSES = ["CANCELLED", "LOST", "DAMAGED"]
ALL_STATUSES = STATUS_FLOW + EXCEPTION_STATUSES
PAID_STATUSES = {"PAID", "LOADED", "IN_TRANSIT", "IN_CUSTOMS", "ARRIVED", "AVAILABLE", "DELIVERED"}

# --- Models ---

class User(BaseModel):
    username: str
    full_name: str
    role: Literal["admin", "director", "agency_manager", "depot_chief", "operator", "accountant", "customer_service", "delivery", "viewer", "supervisor"]
    disabled: Optional[bool] = False
    agency: Optional[str] = None

class UserCreate(User):
    password: str

class Token(BaseModel):
    access_token: Optional[str] = None
    token_type: Optional[str] = None
    role: Optional[str] = None
    full_name: Optional[str] = None
    require_2fa: bool = False
    mfa_token: Optional[str] = None

class TwoFACode(BaseModel):
    code: str

class MFAVerify(BaseModel):
    mfa_token: str
    code: str

class SenderReceiverInfo(BaseModel):
    name: str
    phone: str
    city: str
    address: Optional[str] = None
    country: Optional[str] = None
    email: Optional[str] = None

class Dimensions(BaseModel):
    length: Optional[float] = None
    width: Optional[float] = None
    height: Optional[float] = None

class ParcelCreate(BaseModel):
    direction: Literal["EU_TO_CM", "CM_TO_EU"]
    sender: SenderReceiverInfo
    receiver: SenderReceiverInfo
    content_description: str
    weight_kg: Optional[float] = None
    departure_date: str
    nature: Optional[str] = None
    declared_value: Optional[float] = None
    fragile: Optional[bool] = False
    insured: Optional[bool] = False
    dimensions: Optional[Dimensions] = None
    agency_origin: Optional[str] = None
    agency_destination: Optional[str] = None
    operator: Optional[str] = None

class ParcelUpdate(BaseModel):
    status: Optional[str] = None
    weight_kg: Optional[float] = None
    final_price: Optional[float] = None
    amount_paid: Optional[float] = None
    note: Optional[str] = None
    nature: Optional[str] = None
    declared_value: Optional[float] = None
    fragile: Optional[bool] = None
    insured: Optional[bool] = None
    dimensions: Optional[Dimensions] = None

class NotificationRequest(BaseModel):
    tracking_id: str
    type: Literal["sms", "email", "whatsapp"]
    message: str

# --- Utils ---

def verify_password(plain_password, hashed_password):
    return pwd_context.verify(plain_password, hashed_password)

def get_password_hash(password):
    return pwd_context.hash(password)

def create_access_token(data: dict):
    to_encode = data.copy()
    expire = datetime.now(timezone.utc) + timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)

TOTP_ISSUER = "LOGILINK GLOBAL"
MFA_TOKEN_EXPIRE_MINUTES = 5

def create_mfa_token(username: str):
    expire = datetime.now(timezone.utc) + timedelta(minutes=MFA_TOKEN_EXPIRE_MINUTES)
    return jwt.encode({"sub": username, "typ": "mfa", "exp": expire}, SECRET_KEY, algorithm=ALGORITHM)

def build_qr_data_url(uri: str):
    qr = qrcode.QRCode(box_size=8, border=2)
    qr.add_data(uri)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()

def build_access_response(user: dict) -> dict:
    token = create_access_token(data={"sub": user["username"], "role": user["role"]})
    return {"access_token": token, "token_type": "bearer", "role": user["role"],
            "full_name": user["full_name"], "require_2fa": False}

async def get_current_user(token: str = Depends(oauth2_scheme)):
    credentials_exception = HTTPException(
        status_code=401,
        detail="Could not validate credentials",
        headers={"WWW-Authenticate": "Bearer"},
    )
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        username: str = payload.get("sub")
        if username is None:
            raise credentials_exception
    except JWTError:
        raise credentials_exception
    user = await users_collection.find_one({"username": username}, {"_id": 0})
    if user is None:
        raise credentials_exception
    return user

def require_roles(*allowed_roles: str):
    async def _guard(current_user: dict = Depends(get_current_user)):
        if current_user["role"] not in allowed_roles:
            raise HTTPException(status_code=403, detail="Permissions insuffisantes")
        return current_user
    return _guard

# roles allowed to operate on parcels
OPERATOR_ROLES = ("admin", "director", "agency_manager", "depot_chief", "operator", "supervisor", "customer_service", "accountant")
GLOBAL_ROLES = ("admin", "director")

def agency_query(user: dict) -> dict:
    """Restrict to the user's agency unless they have a global role."""
    if user.get("role") in GLOBAL_ROLES or not user.get("agency"):
        return {}
    ag = user["agency"]
    return {"$or": [{"agency_origin": ag}, {"agency_destination": ag}]}

async def get_pricing() -> dict:
    doc = await settings_collection.find_one({"_key": "pricing"}, {"_id": 0, "_key": 0})
    return doc or {"price_per_kg": {"EU_TO_CM": 8.0, "CM_TO_EU": 10.0}, "vat_percent": 0.0, "currency": "EUR"}

def generate_tracking_id():
    chars = string.ascii_uppercase + string.digits
    suffix = ''.join(random.choice(chars) for _ in range(6))
    return f"LOGI-{suffix}"

def generate_barcode():
    return ''.join(random.choice(string.digits) for _ in range(12))

def now_iso():
    return datetime.now(timezone.utc).isoformat()

async def log_audit(username: str, action: str, target: str, details: str = ""):
    await audit_collection.insert_one({
        "username": username,
        "action": action,
        "target": target,
        "details": details,
        "timestamp": now_iso(),
    })

# --- Seeding ---
@app.on_event("startup")
async def seed_data():
    if await users_collection.count_documents({}) == 0:
        users = [
            {"username": "admin", "full_name": "Administrateur", "role": "admin", "hashed_password": get_password_hash("admin123"), "disabled": False, "agency": None},
            {"username": "operateur", "full_name": "Opérateur Terrain", "role": "operator", "hashed_password": get_password_hash("op123"), "disabled": False, "agency": "AG-LODI"},
            {"username": "superviseur", "full_name": "Superviseur Chef", "role": "supervisor", "hashed_password": get_password_hash("super123"), "disabled": False, "agency": "AG-DOUALA"},
        ]
        await users_collection.insert_many(users)
        print("Users seeded")

    if await agencies_collection.count_documents({}) == 0:
        await agencies_collection.insert_many([
            {"code": "AG-LODI", "name": "Agence Lodi", "country": "Italie", "city": "Lodi", "address": "Via Roma 35, 26866 Lodi", "phone": "+39 3287091255"},
            {"code": "AG-PARIS", "name": "Agence Paris", "country": "France", "city": "Paris", "address": "Paris", "phone": "+33 100000000"},
            {"code": "AG-DOUALA", "name": "Agence Douala", "country": "Cameroun", "city": "Douala", "address": "Akwa, Douala", "phone": "+237 600000000"},
            {"code": "AG-YAOUNDE", "name": "Agence Yaoundé", "country": "Cameroun", "city": "Yaoundé", "address": "Centre, Yaoundé", "phone": "+237 600000001"},
        ])
        print("Agencies seeded")

    if await settings_collection.count_documents({"_key": "pricing"}) == 0:
        await settings_collection.insert_one({
            "_key": "pricing",
            "price_per_kg": {"EU_TO_CM": 8.0, "CM_TO_EU": 10.0},
            "vat_percent": 0.0,
            "currency": "EUR",
        })
        print("Settings seeded")

    # Rich demo parcels
    if await parcels_collection.count_documents({}) < 12:
        await parcels_collection.delete_many({})
        natures = ["Vêtements", "Électronique", "Documents", "Produits alimentaires", "Cosmétiques", "Pièces auto"]
        cities_eu = ["Paris", "Lodi", "Rome", "Bruxelles", "Milan", "Lyon"]
        cities_cm = ["Douala", "Yaoundé", "Bafoussam", "Bamenda", "Kribi", "Garoua"]
        dummy = []
        for i in range(24):
            status = random.choice(ALL_STATUSES)
            direction = random.choice(["EU_TO_CM", "CM_TO_EU"])
            days_ago = random.randint(0, 45)
            created_at = datetime.now(timezone.utc) - timedelta(days=days_ago)
            dep_date = created_at + timedelta(days=2)
            arr_date = dep_date + timedelta(days=7)
            weight = round(random.uniform(1.0, 25.0), 1)
            price = round(weight * random.uniform(6, 12), 2)
            paid = status in PAID_STATUSES
            # build a plausible history up to current status
            history = []
            if status in STATUS_FLOW:
                idx = STATUS_FLOW.index(status)
                for j, s in enumerate(STATUS_FLOW[:idx + 1]):
                    ts = created_at + timedelta(days=j * 0.7)
                    history.append({"status": s, "timestamp": ts.isoformat(), "author": "Opérateur Terrain", "comment": ""})
            else:
                history.append({"status": "REGISTERED", "timestamp": created_at.isoformat(), "author": "Système", "comment": ""})
                history.append({"status": status, "timestamp": (created_at + timedelta(days=3)).isoformat(), "author": "Opérateur Terrain", "comment": "Anomalie signalée"})

            src = cities_eu if direction == "EU_TO_CM" else cities_cm
            dst = cities_cm if direction == "EU_TO_CM" else cities_eu
            dummy.append({
                "tracking_id": generate_tracking_id(),
                "barcode": generate_barcode(),
                "direction": direction,
                "sender": {"name": f"Client Exp. {i}", "phone": f"+33{random.randint(600000000,699999999)}", "city": random.choice(src), "address": "Adresse", "country": "Italie" if direction == "EU_TO_CM" else "Cameroun"},
                "receiver": {"name": f"Client Dest. {i}", "phone": f"+237{random.randint(600000000,699999999)}", "city": random.choice(dst), "address": "Adresse", "country": "Cameroun" if direction == "EU_TO_CM" else "Italie"},
                "content_description": random.choice(natures),
                "nature": random.choice(natures),
                "status": status,
                "created_at": created_at.isoformat(),
                "departure_date": dep_date.strftime("%Y-%m-%d"),
                "estimated_arrival": arr_date.strftime("%Y-%m-%d"),
                "weight_kg": weight,
                "volume_m3": round(random.uniform(0.01, 0.3), 3),
                "dimensions": {"length": random.randint(20, 80), "width": random.randint(20, 60), "height": random.randint(10, 50)},
                "declared_value": round(random.uniform(50, 1200), 2),
                "fragile": random.choice([True, False]),
                "insured": random.choice([True, False]),
                "final_price": price,
                "amount_paid": price if paid else 0.0,
                "operator": "Opérateur Terrain",
                "agency_origin": "AG-LODI" if direction == "EU_TO_CM" else random.choice(["AG-DOUALA", "AG-YAOUNDE"]),
                "agency_destination": random.choice(["AG-DOUALA", "AG-YAOUNDE"]) if direction == "EU_TO_CM" else "AG-LODI",
                "note": "",
                "history": history,
            })
        await parcels_collection.insert_many(dummy)
        print("Parcels seeded")

def generate_pdf_ticket(parcel: dict):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=24, alignment=1, spaceAfter=20)
    normal_style = styles['Normal']

    story.append(Paragraph("TICKET D'ENVOI - LOGILINK GLOBAL", title_style))
    story.append(Spacer(1, 0.3 * inch))
    story.append(Paragraph(f"<b>N° SUIVI: {parcel['tracking_id']}</b>", ParagraphStyle('Tracking', parent=styles['Heading2'], fontSize=28, alignment=1, textColor=colors.HexColor('#0F172A'))))
    story.append(Spacer(1, 0.2 * inch))

    dir_text = "EUROPE -> CAMEROUN" if parcel['direction'] == "EU_TO_CM" else "CAMEROUN -> EUROPE"
    story.append(Paragraph(f"DIRECTION: {dir_text}", ParagraphStyle('Dir', parent=styles['Normal'], fontSize=14, alignment=1)))
    story.append(Spacer(1, 0.3 * inch))

    qr = qrcode.QRCode(box_size=10, border=4)
    qr.add_data(f"https://logilink.com/track?id={parcel['tracking_id']}")
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    img_buffer = io.BytesIO()
    img.save(img_buffer)
    img_buffer.seek(0)
    story.append(RLImage(img_buffer, width=1.8 * inch, height=1.8 * inch))
    story.append(Spacer(1, 0.2 * inch))

    # Code 128 barcode
    bc_value = parcel.get('barcode') or parcel['tracking_id']
    barcode = code128.Code128(bc_value, barHeight=0.5 * inch, barWidth=1.1)
    story.append(barcode)
    story.append(Paragraph(bc_value, ParagraphStyle('bc', parent=styles['Normal'], fontSize=9, alignment=1)))
    story.append(Spacer(1, 0.3 * inch))

    data = [
        ["EXPÉDITEUR", "DESTINATAIRE"],
        [f"{parcel['sender']['name']}\n{parcel['sender']['phone']}\n{parcel['sender']['city']}",
         f"{parcel['receiver']['name']}\n{parcel['receiver']['phone']}\n{parcel['receiver']['city']}"],
        ["CONTENU", parcel['content_description']],
        ["DATE DÉPART PRÉVUE", parcel['departure_date']]
    ]
    t = Table(data, colWidths=[3 * inch, 3 * inch])
    t.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('BACKGROUND', (0, 0), (1, 0), colors.HexColor('#F1F5F9')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('PADDING', (0, 0), (-1, -1), 12),
        ('FONTSIZE', (0, 0), (-1, -1), 12),
        ('FONTNAME', (0, 0), (1, 0), 'Helvetica-Bold'),
        ('SPAN', (0, 2), (1, 2)),
        ('SPAN', (0, 3), (1, 3)),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.8 * inch))
    story.append(Paragraph("Merci de coller ce ticket sur votre colis avant le dépôt.", normal_style))
    doc.build(story)
    buffer.seek(0)
    return buffer

def generate_invoice_pdf(inv: dict):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph("FACTURE - LOGILINK GLOBAL", ParagraphStyle('T', parent=styles['Heading1'], fontSize=22, alignment=0, textColor=colors.HexColor('#0F172A'))))
    story.append(Paragraph(f"N° {inv['invoice_number']} · {inv['created_at'][:10]}", ParagraphStyle('S', parent=styles['Normal'], fontSize=10, textColor=colors.grey)))
    story.append(Spacer(1, 0.3 * inch))
    story.append(Paragraph(f"Colis: <b>{inv['tracking_id']}</b>", styles['Normal']))
    story.append(Paragraph(f"Client: {inv['client'].get('name','')} — {inv['client'].get('phone','')}", styles['Normal']))
    story.append(Spacer(1, 0.3 * inch))
    cur = inv.get('currency', 'EUR')
    rows = [
        ["Description", "Poids", "Prix/kg", "Montant"],
        [f"Transport {inv['direction'].replace('_',' ')}", f"{inv['weight_kg']} kg", f"{inv['unit_price']} {cur}", f"{inv['subtotal']} {cur}"],
    ]
    t = Table(rows, colWidths=[2.8 * inch, 1 * inch, 1 * inch, 1.2 * inch])
    t.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('PADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(t)
    story.append(Spacer(1, 0.2 * inch))
    totals = [
        ["Sous-total", f"{inv['subtotal']} {cur}"],
        [f"Remise ({inv['discount_percent']}%)", f"-{inv['discount']} {cur}"],
        [f"TVA ({inv['vat_percent']}%)", f"{inv['vat']} {cur}"],
        ["TOTAL", f"{inv['total']} {cur}"],
    ]
    tt = Table(totals, colWidths=[4 * inch, 2 * inch])
    tt.setStyle(TableStyle([
        ('FONTSIZE', (0, 0), (-1, -1), 11),
        ('PADDING', (0, 0), (-1, -1), 6),
        ('LINEABOVE', (0, 3), (-1, 3), 1, colors.HexColor('#0F172A')),
        ('FONTNAME', (0, 3), (-1, 3), 'Helvetica-Bold'),
        ('TEXTCOLOR', (1, 3), (1, 3), colors.HexColor('#EA580C')),
    ]))
    story.append(tt)
    story.append(Spacer(1, 0.4 * inch))
    status_txt = {"paid": "PAYÉE", "partial": "PARTIELLE", "unpaid": "NON PAYÉE"}.get(inv.get('status'), '')
    story.append(Paragraph(f"Statut: <b>{status_txt}</b> — Payé: {inv.get('amount_paid',0)} {cur}", styles['Normal']))
    story.append(Spacer(1, 0.5 * inch))
    story.append(Paragraph("Merci de votre confiance — LOGILINK GLOBAL", ParagraphStyle('f', parent=styles['Normal'], fontSize=9, textColor=colors.grey)))
    doc.build(story)
    buffer.seek(0)
    return buffer

def generate_report_excel(period, parcels, summary):
    wb = Workbook()
    ws = wb.active
    ws.title = "Résumé"
    ws.append(["Rapport LOGILINK GLOBAL", period])
    ws.append(["Généré le", datetime.now().strftime("%d/%m/%Y %H:%M")])
    ws.append([])
    ws.append(["Total colis", summary["total"]])
    ws.append(["Encaissé (€)", summary["revenue_collected"]])
    ws.append(["Reste à encaisser (€)", summary["revenue_pending"]])
    ws.append(["Poids total (kg)", summary["total_weight"]])
    ws.append([])
    ws.append(["Statut", "Nombre"])
    for k, v in summary["status_counts"].items():
        ws.append([k, v])
    ws2 = wb.create_sheet("Colis")
    ws2.append(["N° Suivi", "Direction", "Statut", "Expéditeur", "Destinataire", "Ville dép.", "Ville arr.", "Poids (kg)", "Prix (€)", "Créé le"])
    for p in parcels:
        ws2.append([p.get("tracking_id"), p.get("direction"), p.get("status"),
                    p.get("sender", {}).get("name"), p.get("receiver", {}).get("name"),
                    p.get("sender", {}).get("city"), p.get("receiver", {}).get("city"),
                    p.get("weight_kg"), p.get("final_price"), (p.get("created_at") or "")[:10]])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

def generate_report_pdf(period, parcels, summary):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph(f"RAPPORT {period.upper()} — LOGILINK GLOBAL", ParagraphStyle('T', parent=styles['Heading1'], fontSize=20, textColor=colors.HexColor('#0F172A'))))
    story.append(Paragraph(datetime.now().strftime("%d/%m/%Y %H:%M"), styles['Normal']))
    story.append(Spacer(1, 0.3 * inch))
    srows = [["Total colis", str(summary["total"])], ["Encaissé", f"{summary['revenue_collected']} €"],
             ["Reste à encaisser", f"{summary['revenue_pending']} €"], ["Poids total", f"{summary['total_weight']} kg"]]
    st = Table(srows, colWidths=[3 * inch, 2 * inch])
    st.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')), ('FONTSIZE', (0, 0), (-1, -1), 11), ('PADDING', (0, 0), (-1, -1), 6)]))
    story.append(st)
    story.append(Spacer(1, 0.3 * inch))
    story.append(Paragraph("Colis de la période", styles['Heading2']))
    data = [["N° Suivi", "Statut", "Trajet", "Prix"]]
    for p in parcels[:40]:
        data.append([p.get("tracking_id"), p.get("status"),
                     f"{p.get('sender', {}).get('city', '')} > {p.get('receiver', {}).get('city', '')}",
                     f"{p.get('final_price', 0)} €"])
    t = Table(data, colWidths=[1.6 * inch, 1.4 * inch, 2.2 * inch, 1 * inch])
    t.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
                           ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
                           ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('FONTSIZE', (0, 0), (-1, -1), 8), ('PADDING', (0, 0), (-1, -1), 5)]))
    story.append(t)
    doc.build(story)
    buffer.seek(0)
    return buffer

# --- Auth Routes ---

@api_router.post("/auth/login", response_model=Token)
async def login(form_data: OAuth2PasswordRequestForm = Depends()):
    user = await users_collection.find_one({"username": form_data.username})
    if not user or not verify_password(form_data.password, user["hashed_password"]):
        raise HTTPException(status_code=400, detail="Identifiant ou mot de passe incorrect")
    if user.get("disabled"):
        raise HTTPException(status_code=403, detail="Compte désactivé")
    if user.get("twofa", {}).get("enabled"):
        return {"require_2fa": True, "mfa_token": create_mfa_token(user["username"])}
    return build_access_response(user)

@api_router.post("/auth/2fa/verify", response_model=Token)
async def verify_2fa_login(payload: MFAVerify):
    try:
        claims = jwt.decode(payload.mfa_token, SECRET_KEY, algorithms=[ALGORITHM])
    except JWTError:
        raise HTTPException(status_code=401, detail="Session expirée, reconnectez-vous")
    if claims.get("typ") != "mfa":
        raise HTTPException(status_code=401, detail="Jeton invalide")
    user = await users_collection.find_one({"username": claims.get("sub")})
    if not user or not user.get("twofa", {}).get("enabled"):
        raise HTTPException(status_code=401, detail="2FA non activé")
    totp = pyotp.TOTP(user["twofa"]["secret"])
    if not totp.verify(payload.code, valid_window=1):
        raise HTTPException(status_code=401, detail="Code incorrect")
    return build_access_response(user)

@api_router.get("/auth/me")
async def read_me(current_user: dict = Depends(get_current_user)):
    return {"username": current_user["username"], "full_name": current_user["full_name"],
            "role": current_user["role"], "twofa_enabled": current_user.get("twofa", {}).get("enabled", False)}

@api_router.post("/auth/2fa/setup")
async def setup_2fa(current_user: dict = Depends(get_current_user)):
    secret = pyotp.random_base32()
    await users_collection.update_one({"username": current_user["username"]},
                                      {"$set": {"twofa.pending_secret": secret, "twofa.enabled": current_user.get("twofa", {}).get("enabled", False)}})
    uri = pyotp.TOTP(secret).provisioning_uri(name=current_user["username"], issuer_name=TOTP_ISSUER)
    return {"provisioning_uri": uri, "qr_data_url": build_qr_data_url(uri)}

@api_router.post("/auth/2fa/enable")
async def enable_2fa(payload: TwoFACode, current_user: dict = Depends(get_current_user)):
    pending = current_user.get("twofa", {}).get("pending_secret")
    if not pending:
        raise HTTPException(status_code=400, detail="Aucune configuration 2FA en cours")
    if not pyotp.TOTP(pending).verify(payload.code, valid_window=1):
        raise HTTPException(status_code=400, detail="Code incorrect")
    await users_collection.update_one({"username": current_user["username"]},
                                      {"$set": {"twofa.enabled": True, "twofa.secret": pending},
                                       "$unset": {"twofa.pending_secret": ""}})
    await log_audit(current_user["username"], "ENABLE_2FA", current_user["username"])
    return {"ok": True}

@api_router.post("/auth/2fa/disable")
async def disable_2fa(payload: TwoFACode, current_user: dict = Depends(get_current_user)):
    secret = current_user.get("twofa", {}).get("secret")
    if not secret:
        raise HTTPException(status_code=400, detail="2FA non activé")
    if not pyotp.TOTP(secret).verify(payload.code, valid_window=1):
        raise HTTPException(status_code=400, detail="Code incorrect")
    await users_collection.update_one({"username": current_user["username"]},
                                      {"$set": {"twofa.enabled": False}, "$unset": {"twofa.secret": "", "twofa.pending_secret": ""}})
    await log_audit(current_user["username"], "DISABLE_2FA", current_user["username"])
    return {"ok": True}

# --- Users (admin) ---

@api_router.get("/users", response_model=List[User])
async def list_users(current_user: dict = Depends(require_roles("admin", "director"))):
    return await users_collection.find({}, {"_id": 0, "hashed_password": 0}).to_list(200)

@api_router.post("/users", response_model=User)
async def create_user(user_in: UserCreate, current_user: dict = Depends(require_roles("admin", "director"))):
    if await users_collection.find_one({"username": user_in.username}):
        raise HTTPException(status_code=400, detail="Nom d'utilisateur déjà pris")
    user_doc = user_in.dict()
    user_doc["hashed_password"] = get_password_hash(user_in.password)
    del user_doc["password"]
    await users_collection.insert_one(user_doc)
    await log_audit(current_user["username"], "CREATE_USER", user_in.username)
    return user_in

# --- Parcels ---

@api_router.post("/parcels")
async def create_parcel(parcel_in: ParcelCreate):
    tracking_id = generate_tracking_id()
    dep_date = datetime.strptime(parcel_in.departure_date, "%Y-%m-%d")
    arrival_date = dep_date + timedelta(days=7)
    parcel_doc = parcel_in.dict()
    author = parcel_in.operator or "Client (Web)"
    parcel_doc.update({
        "tracking_id": tracking_id,
        "barcode": generate_barcode(),
        "status": "REGISTERED",
        "created_at": now_iso(),
        "estimated_arrival": arrival_date.strftime("%Y-%m-%d"),
        "weight_kg": parcel_in.weight_kg or 0.0,
        "volume_m3": 0.0,
        "final_price": 0.0,
        "amount_paid": 0.0,
        "note": "",
        "history": [{"status": "REGISTERED", "timestamp": now_iso(), "author": author, "comment": "Colis enregistré"}],
    })
    await parcels_collection.insert_one(parcel_doc)
    parcel_doc.pop("_id", None)
    return parcel_doc

@api_router.get("/parcels/{tracking_id}")
async def get_parcel(tracking_id: str):
    parcel = await parcels_collection.find_one({"tracking_id": tracking_id}, {"_id": 0})
    if not parcel:
        raise HTTPException(status_code=404, detail="Colis non trouvé")
    return parcel

@api_router.get("/parcels/{tracking_id}/pdf")
async def get_parcel_pdf(tracking_id: str):
    parcel = await parcels_collection.find_one({"tracking_id": tracking_id}, {"_id": 0})
    if not parcel:
        raise HTTPException(status_code=404, detail="Colis non trouvé")
    pdf_buffer = generate_pdf_ticket(parcel)
    return StreamingResponse(pdf_buffer, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename=ticket-{tracking_id}.pdf"})

@api_router.get("/parcels")
async def list_parcels(current_user: dict = Depends(require_roles(*OPERATOR_ROLES))):
    return await parcels_collection.find(agency_query(current_user), {"_id": 0}).sort("created_at", -1).to_list(2000)

@api_router.patch("/parcels/{tracking_id}")
async def update_parcel(tracking_id: str, update: ParcelUpdate, current_user: dict = Depends(require_roles(*OPERATOR_ROLES))):
    update_data = {k: v for k, v in update.dict().items() if v is not None}
    if not update_data:
        return {"message": "No data"}
    parcel = await parcels_collection.find_one({"tracking_id": tracking_id})
    if not parcel:
        raise HTTPException(status_code=404, detail="Colis non trouvé")
    history_entry = None
    if "status" in update_data and update_data["status"] != parcel.get("status"):
        history_entry = {"status": update_data["status"], "timestamp": now_iso(),
                         "author": current_user["full_name"], "comment": update_data.get("note", "")}
    ops = {"$set": update_data}
    if history_entry:
        ops["$push"] = {"history": history_entry}
    await parcels_collection.update_one({"tracking_id": tracking_id}, ops)
    await log_audit(current_user["username"], "UPDATE_PARCEL", tracking_id, str(update_data))
    return {"message": "Parcel updated", "updated_fields": update_data}

@api_router.patch("/parcels/{tracking_id}/status")
async def update_status(tracking_id: str, status: str, current_user: dict = Depends(require_roles(*OPERATOR_ROLES))):
    parcel = await parcels_collection.find_one({"tracking_id": tracking_id})
    if not parcel:
        raise HTTPException(status_code=404, detail="Colis non trouvé")
    entry = {"status": status, "timestamp": now_iso(), "author": current_user["full_name"], "comment": ""}
    await parcels_collection.update_one({"tracking_id": tracking_id},
                                        {"$set": {"status": status}, "$push": {"history": entry}})
    await log_audit(current_user["username"], "UPDATE_STATUS", tracking_id, status)
    return {"message": "Status updated"}

# --- Schedule ---

@api_router.get("/schedule")
async def get_schedule():
    today = datetime.now()
    schedule = {"eu_to_cm": [], "cm_to_eu": []}
    next_friday = today + timedelta((4 - today.weekday()) % 7)
    next_saturday = today + timedelta((5 - today.weekday()) % 7)
    for i in range(4):
        schedule["eu_to_cm"].append((next_friday + timedelta(weeks=i)).strftime("%Y-%m-%d"))
        schedule["cm_to_eu"].append((next_saturday + timedelta(weeks=i)).strftime("%Y-%m-%d"))
    return schedule

# --- Stats / Dashboard ---

@api_router.get("/stats")
async def get_stats(current_user: dict = Depends(require_roles(*OPERATOR_ROLES))):
    parcels = await parcels_collection.find(agency_query(current_user), {"_id": 0}).to_list(5000)
    total = len(parcels)
    status_counts = {s: 0 for s in ALL_STATUSES}
    revenue_collected = 0.0
    revenue_pending = 0.0
    total_weight = 0.0
    total_volume = 0.0
    senders, receivers, clients = set(), set(), set()
    dir_counts = {"EU_TO_CM": 0, "CM_TO_EU": 0}
    delivery_days = []
    daily = defaultdict(int)
    monthly = defaultdict(int)

    for p in parcels:
        st = p.get("status", "REGISTERED")
        status_counts[st] = status_counts.get(st, 0) + 1
        price = float(p.get("final_price") or 0)
        paid = float(p.get("amount_paid") or 0)
        revenue_collected += paid
        revenue_pending += max(price - paid, 0)
        total_weight += float(p.get("weight_kg") or 0)
        total_volume += float(p.get("volume_m3") or 0)
        s_ph = p.get("sender", {}).get("phone")
        r_ph = p.get("receiver", {}).get("phone")
        if s_ph:
            senders.add(s_ph); clients.add(s_ph)
        if r_ph:
            receivers.add(r_ph); clients.add(r_ph)
        dir_counts[p.get("direction", "EU_TO_CM")] = dir_counts.get(p.get("direction", "EU_TO_CM"), 0) + 1
        # activity by created day
        try:
            created = datetime.fromisoformat(p["created_at"].replace("Z", "+00:00"))
            daily[created.strftime("%Y-%m-%d")] += 1
            monthly[created.strftime("%Y-%m")] += 1
        except Exception:
            pass
        # delivery time
        if st == "DELIVERED":
            hist = p.get("history", [])
            delivered_at = next((h["timestamp"] for h in hist if h["status"] == "DELIVERED"), None)
            try:
                c = datetime.fromisoformat(p["created_at"].replace("Z", "+00:00"))
                d = datetime.fromisoformat(delivered_at.replace("Z", "+00:00")) if delivered_at else None
                if d:
                    delivery_days.append((d - c).days)
            except Exception:
                pass

    delivered = status_counts.get("DELIVERED", 0)
    incidents = status_counts.get("LOST", 0) + status_counts.get("DAMAGED", 0) + status_counts.get("CANCELLED", 0)
    in_transit = sum(status_counts.get(s, 0) for s in ("LOADED", "IN_TRANSIT", "IN_CUSTOMS"))
    registered = status_counts.get("REGISTERED", 0) + status_counts.get("CREATED", 0)
    waiting = sum(status_counts.get(s, 0) for s in ("RECEIVED_AT_DEPOT", "CONTROLLED", "WEIGHED", "PACKED", "INVOICED", "PAID"))
    arrived = sum(status_counts.get(s, 0) for s in ("ARRIVED", "AVAILABLE"))

    # daily activity last 14 days
    daily_series = []
    for i in range(13, -1, -1):
        d = (datetime.now() - timedelta(days=i)).strftime("%Y-%m-%d")
        daily_series.append({"date": d, "count": daily.get(d, 0)})
    # monthly last 6 months
    monthly_series = []
    for i in range(5, -1, -1):
        d = (datetime.now() - timedelta(days=i * 30)).strftime("%Y-%m")
        monthly_series.append({"month": d, "count": monthly.get(d, 0)})

    success_rate = round((delivered / total) * 100, 1) if total else 0
    avg_delivery = round(sum(delivery_days) / len(delivery_days), 1) if delivery_days else 0

    return {
        "total": total,
        "registered": registered,
        "waiting": waiting,
        "transit": in_transit,
        "arrived": arrived,
        "delivered": delivered,
        "incidents": incidents,
        "delayed": 0,
        "status_counts": status_counts,
        "revenue_collected": round(revenue_collected, 2),
        "revenue_pending": round(revenue_pending, 2),
        "total_weight": round(total_weight, 1),
        "total_volume": round(total_volume, 3),
        "clients": len(clients),
        "senders": len(senders),
        "receivers": len(receivers),
        "direction_counts": dir_counts,
        "success_rate": success_rate,
        "avg_delivery_days": avg_delivery,
        "daily_activity": daily_series,
        "monthly_activity": monthly_series,
    }

@api_router.get("/audit")
async def list_audit(current_user: dict = Depends(require_roles("admin", "director"))):
    return await audit_collection.find({}, {"_id": 0}).sort("timestamp", -1).to_list(100)

@api_router.post("/notify/simulate")
async def simulate_notification(req: NotificationRequest, current_user: dict = Depends(require_roles(*OPERATOR_ROLES))):
    await asyncio.sleep(0.5)
    return {"success": True, "message": f"SIMULATION: {req.type.upper()} envoyé pour {req.tracking_id}", "details": req.message}

# ---------- Agencies ----------
class AgencyModel(BaseModel):
    code: str
    name: str
    country: str
    city: str
    address: Optional[str] = None
    phone: Optional[str] = None

@api_router.get("/agencies")
async def list_agencies(current_user: dict = Depends(get_current_user)):
    return await agencies_collection.find({}, {"_id": 0}).to_list(200)

@api_router.post("/agencies")
async def create_agency(a: AgencyModel, current_user: dict = Depends(require_roles("admin", "director"))):
    if await agencies_collection.find_one({"code": a.code}):
        raise HTTPException(status_code=400, detail="Code agence déjà utilisé")
    await agencies_collection.insert_one(a.dict())
    await log_audit(current_user["username"], "CREATE_AGENCY", a.code)
    return a

@api_router.delete("/agencies/{code}")
async def delete_agency(code: str, current_user: dict = Depends(require_roles("admin", "director"))):
    await agencies_collection.delete_one({"code": code})
    await log_audit(current_user["username"], "DELETE_AGENCY", code)
    return {"message": "deleted"}

# ---------- Settings / pricing ----------
class PricingUpdate(BaseModel):
    price_per_kg: Dict[str, float]
    vat_percent: float
    currency: str = "EUR"

@api_router.get("/settings")
async def read_settings(current_user: dict = Depends(require_roles(*OPERATOR_ROLES))):
    return await get_pricing()

@api_router.put("/settings")
async def write_settings(p: PricingUpdate, current_user: dict = Depends(require_roles("admin", "director"))):
    await settings_collection.update_one({"_key": "pricing"}, {"$set": {**p.dict(), "_key": "pricing"}}, upsert=True)
    await log_audit(current_user["username"], "UPDATE_SETTINGS", "pricing")
    return await get_pricing()

# ---------- Clients (derived from parcels) ----------
@api_router.get("/clients")
async def list_clients(current_user: dict = Depends(require_roles(*OPERATOR_ROLES))):
    parcels = await parcels_collection.find(agency_query(current_user), {"_id": 0}).to_list(5000)
    clients = {}
    for p in parcels:
        for role_key in ("sender", "receiver"):
            info = p.get(role_key) or {}
            ph = info.get("phone")
            if not ph:
                continue
            c = clients.setdefault(ph, {"phone": ph, "name": info.get("name"), "city": info.get("city"),
                                        "country": info.get("country"), "parcels": 0, "total_spent": 0.0,
                                        "total_weight": 0.0, "as_sender": 0, "as_receiver": 0, "last_shipment": None})
            c["parcels"] += 1
            c["total_weight"] += float(p.get("weight_kg") or 0)
            c["total_spent"] += float(p.get("final_price") or 0)
            c["as_sender" if role_key == "sender" else "as_receiver"] += 1
            ca = p.get("created_at")
            if ca and (not c["last_shipment"] or ca > c["last_shipment"]):
                c["last_shipment"] = ca
            if not c["name"]:
                c["name"] = info.get("name")
    out = list(clients.values())
    out.sort(key=lambda x: x["parcels"], reverse=True)
    for c in out:
        c["total_spent"] = round(c["total_spent"], 2)
        c["total_weight"] = round(c["total_weight"], 1)
    return out

@api_router.get("/clients/{phone}")
async def client_detail(phone: str, current_user: dict = Depends(require_roles(*OPERATOR_ROLES))):
    q = agency_query(current_user)
    base = {"$or": [{"sender.phone": phone}, {"receiver.phone": phone}]}
    if q:
        base = {"$and": [q, base]}
    parcels = await parcels_collection.find(base, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return {"phone": phone, "parcels": parcels}

# ---------- Invoices ----------
class InvoiceCreate(BaseModel):
    tracking_id: str
    discount_percent: Optional[float] = 0.0
    unit_price: Optional[float] = None

def gen_invoice_number():
    return "INV-" + ''.join(random.choice(string.digits) for _ in range(6))

@api_router.get("/invoices")
async def list_invoices(current_user: dict = Depends(require_roles(*OPERATOR_ROLES))):
    return await invoices_collection.find({}, {"_id": 0}).sort("created_at", -1).to_list(1000)

@api_router.post("/invoices")
async def create_invoice(inv: InvoiceCreate, current_user: dict = Depends(require_roles("admin", "director", "accountant", "agency_manager"))):
    parcel = await parcels_collection.find_one({"tracking_id": inv.tracking_id}, {"_id": 0})
    if not parcel:
        raise HTTPException(status_code=404, detail="Colis non trouvé")
    if await invoices_collection.find_one({"tracking_id": inv.tracking_id}):
        raise HTTPException(status_code=400, detail="Facture déjà existante pour ce colis")
    pricing = await get_pricing()
    weight = float(parcel.get("weight_kg") or 0)
    if weight <= 0:
        raise HTTPException(status_code=400, detail="Poids du colis requis pour facturer")
    unit = inv.unit_price if inv.unit_price is not None else pricing["price_per_kg"].get(parcel["direction"], 8.0)
    subtotal = round(weight * unit, 2)
    discount = round(subtotal * (inv.discount_percent or 0) / 100, 2)
    taxable = subtotal - discount
    vat = round(taxable * pricing["vat_percent"] / 100, 2)
    total = round(taxable + vat, 2)
    number = gen_invoice_number()
    doc = {
        "invoice_number": number, "tracking_id": inv.tracking_id, "direction": parcel["direction"],
        "client": parcel.get("sender"), "receiver": parcel.get("receiver"),
        "weight_kg": weight, "unit_price": unit, "subtotal": subtotal,
        "discount_percent": inv.discount_percent or 0, "discount": discount,
        "vat_percent": pricing["vat_percent"], "vat": vat, "total": total,
        "currency": pricing["currency"], "amount_paid": 0.0, "status": "unpaid",
        "created_at": now_iso(), "created_by": current_user["full_name"],
    }
    await invoices_collection.insert_one(doc)
    await parcels_collection.update_one({"tracking_id": inv.tracking_id},
                                        {"$set": {"final_price": total, "status": "INVOICED"},
                                         "$push": {"history": {"status": "INVOICED", "timestamp": now_iso(), "author": current_user["full_name"], "comment": f"Facture {number}"}}})
    await log_audit(current_user["username"], "CREATE_INVOICE", number, inv.tracking_id)
    doc.pop("_id", None)
    return doc

@api_router.patch("/invoices/{number}/pay")
async def pay_invoice(number: str, amount: float, current_user: dict = Depends(require_roles("admin", "director", "accountant", "agency_manager"))):
    inv = await invoices_collection.find_one({"invoice_number": number})
    if not inv:
        raise HTTPException(status_code=404, detail="Facture non trouvée")
    paid = round(float(inv.get("amount_paid", 0)) + amount, 2)
    status = "paid" if paid >= inv["total"] else ("partial" if paid > 0 else "unpaid")
    await invoices_collection.update_one({"invoice_number": number}, {"$set": {"amount_paid": paid, "status": status}})
    await parcels_collection.update_one({"tracking_id": inv["tracking_id"]}, {"$set": {"amount_paid": paid}})
    if status == "paid":
        await parcels_collection.update_one({"tracking_id": inv["tracking_id"], "status": "INVOICED"},
                                            {"$set": {"status": "PAID"}, "$push": {"history": {"status": "PAID", "timestamp": now_iso(), "author": current_user["full_name"], "comment": f"Paiement facture {number}"}}})
    await log_audit(current_user["username"], "PAY_INVOICE", number, str(amount))
    return {"invoice_number": number, "amount_paid": paid, "status": status}

@api_router.get("/invoices/{number}/pdf")
async def invoice_pdf(number: str):
    inv = await invoices_collection.find_one({"invoice_number": number}, {"_id": 0})
    if not inv:
        raise HTTPException(status_code=404, detail="Facture non trouvée")
    buf = generate_invoice_pdf(inv)
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename={number}.pdf"})

# ---------- E-commerce integration (generic inbound webhook) ----------
class DefaultParty(BaseModel):
    name: str = ""
    phone: str = ""
    city: str = ""
    country: Optional[str] = None

class ApiKeyCreate(BaseModel):
    label: str
    platform: Literal["generic", "shopify", "woocommerce"] = "generic"
    default_direction: Literal["EU_TO_CM", "CM_TO_EU"] = "EU_TO_CM"
    default_sender: Optional[DefaultParty] = None

class EcomShipment(BaseModel):
    direction: Literal["EU_TO_CM", "CM_TO_EU"]
    sender: SenderReceiverInfo
    receiver: SenderReceiverInfo
    content_description: str
    weight_kg: Optional[float] = None
    external_order_id: Optional[str] = None
    departure_date: Optional[str] = None

@api_router.get("/integrations/keys")
async def list_keys(current_user: dict = Depends(require_roles("admin", "director"))):
    return await apikeys_collection.find({}, {"_id": 0}).to_list(100)

@api_router.post("/integrations/keys")
async def create_key(body: ApiKeyCreate, current_user: dict = Depends(require_roles("admin", "director"))):
    key = "sk_live_" + secrets.token_hex(16)
    doc = {"key": key, "label": body.label, "platform": body.platform,
           "default_direction": body.default_direction,
           "default_sender": body.default_sender.dict() if body.default_sender else None,
           "created_at": now_iso(), "created_by": current_user["username"], "active": True}
    await apikeys_collection.insert_one(doc)
    await log_audit(current_user["username"], "CREATE_APIKEY", body.label)
    doc.pop("_id", None)
    return doc

@api_router.delete("/integrations/keys/{key}")
async def revoke_key(key: str, current_user: dict = Depends(require_roles("admin", "director"))):
    await apikeys_collection.update_one({"key": key}, {"$set": {"active": False}})
    await log_audit(current_user["username"], "REVOKE_APIKEY", key[:16])
    return {"message": "revoked"}

async def verify_api_key(x_api_key: Optional[str] = Header(None)):
    if not x_api_key:
        raise HTTPException(status_code=401, detail="X-API-Key requis")
    doc = await apikeys_collection.find_one({"key": x_api_key, "active": True})
    if not doc:
        raise HTTPException(status_code=403, detail="Clé API invalide ou révoquée")
    return doc

@api_router.post("/integrations/shipments")
async def ecom_create_shipment(order: EcomShipment, api_key: dict = Depends(verify_api_key)):
    tracking_id = generate_tracking_id()
    dep = order.departure_date or (datetime.now() + timedelta(days=2)).strftime("%Y-%m-%d")
    arrival = (datetime.strptime(dep, "%Y-%m-%d") + timedelta(days=7)).strftime("%Y-%m-%d")
    doc = order.dict()
    doc.update({
        "tracking_id": tracking_id, "barcode": generate_barcode(), "status": "REGISTERED",
        "created_at": now_iso(), "departure_date": dep, "estimated_arrival": arrival,
        "weight_kg": order.weight_kg or 0.0, "volume_m3": 0.0, "final_price": 0.0, "amount_paid": 0.0,
        "note": "", "source": "ecommerce", "external_order_id": order.external_order_id,
        "agency_origin": "AG-LODI" if order.direction == "EU_TO_CM" else "AG-DOUALA",
        "agency_destination": "AG-DOUALA" if order.direction == "EU_TO_CM" else "AG-LODI",
        "history": [{"status": "REGISTERED", "timestamp": now_iso(), "author": f"E-commerce ({api_key['label']})", "comment": "Commande importée automatiquement"}],
    })
    await parcels_collection.insert_one(doc)
    return {"tracking_id": tracking_id, "tracking_url": f"/track?id={tracking_id}", "status": "REGISTERED"}

app.include_router(api_router)
